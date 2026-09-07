# -*- coding: utf-8 -*-
"""
02_maj_classeur_avec_deflation.py -- Reconstruit le classeur final en
ajoutant un 8e bloc "SELECTION APRES DEFLATION" (sarcelle), uniquement
pour les branches concernees (Electricite, Finances, Immobilier).
"""
import pickle
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DOSSIER = '/home/claude/phase2'
FILL_CIBLE  = PatternFill("solid", fgColor="548235")
FILL_ANCRE  = PatternFill("solid", fgColor="BF9000")
FILL_SELECT = PatternFill("solid", fgColor="7030A0")
FILL_DEFLATE= PatternFill("solid", fgColor="00B0A0")   # sarcelle
FILL_T      = PatternFill("solid", fgColor="1F4E78")
FILL_M      = PatternFill("solid", fgColor="C55A11")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SEP = 2

wb = Workbook()
wb.remove(wb.active)

def construire_dict(items):
    d = {}
    for ind in items:
        nom = ind['nom']; base = nom; i = 2
        while nom in d:
            nom = f"{base} ({i})"; i += 1
        d[nom] = dict(ind['serie'])
    return d

def ecrire_bloc(ws, col_depart, titre, fill, series_dict):
    if not series_dict: return col_depart
    df = pd.DataFrame(series_dict).sort_index()
    n_col = len(df.columns)
    c1 = get_column_letter(col_depart); c2 = get_column_letter(col_depart + n_col)
    ws.merge_cells(f"{c1}1:{c2}1")
    cell = ws.cell(row=1, column=col_depart, value=titre)
    cell.font = Font(size=10, bold=True, color="FFFFFF"); cell.fill = fill
    cell.alignment = Alignment(vertical="center", horizontal="center")
    cell2 = ws.cell(row=3, column=col_depart, value="Date")
    cell2.font = Font(bold=True, color="FFFFFF"); cell2.fill = fill; cell2.border = BORDER
    ws.column_dimensions[get_column_letter(col_depart)].width = 12
    for j, col in enumerate(df.columns):
        cc = col_depart + 1 + j
        c = ws.cell(row=3, column=cc, value=col)
        c.font = Font(bold=True, size=9, color="FFFFFF"); c.fill = fill
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(cc)].width = 16
    ws.row_dimensions[3].height = 55
    for i, (d, row) in enumerate(df.iterrows()):
        r = 4 + i
        c0 = ws.cell(row=r, column=col_depart, value=d)
        c0.number_format = "YYYY-MM-DD"; c0.border = BORDER; c0.font = Font(size=9)
        for j, col in enumerate(df.columns):
            cc = col_depart + 1 + j; val = row[col]
            c = ws.cell(row=r, column=cc, value=round(float(val),4) if pd.notna(val) else None)
            if pd.notna(val): c.number_format = "#,##0.0000"
            c.border = BORDER; c.font = Font(size=9)
    return col_depart + 1 + n_col + SEP

with open(f'{DOSSIER}/resultats/niveau2_resultats.pkl','rb') as f:
    res = pickle.load(f)
with open(f'{DOSSIER}/resultats/niveau2_immobilier.pkl','rb') as f:
    res_immo = pickle.load(f)
with open('nouvelles_series_deflatees.pkl','rb') as f:
    nouvelles = pickle.load(f)

BRANCHES_STANDARD = [
    ('Agriculture', 'agriculture'), ('Peche', 'peche'),
    ('Industrie transformation', 'industrie_transfo'), ('Industrie extraction', 'ind_extraction'),
    ('Finances assurances', 'finances'), ('Hebergement restauration', 'hebergement'),
    ('Construction', 'construction'), ('Commerce', 'commerce'),
    ('Transports', 'transports'), ('Electricite gaz eau', 'electricite'),
    ('Information communication', 'infocom'),
]
recap = []

def traiter_branche(nom_feuille, prefix, ancres_objs, res_reste, brut_cible):
    ancres_trim = [a for a in ancres_objs if a['freq']=='trimestriel']
    ancres_mens = [a for a in ancres_objs if a['freq']=='mensuel']
    selectionnes = [r['obj'] for r in res_reste if r['retenu']]
    non_selectionnes = [r['obj'] for r in res_reste if not r['retenu']]
    sel_trim = [s for s in selectionnes if s['freq']=='trimestriel']
    sel_mens = [s for s in selectionnes if s['freq']=='mensuel']
    reste_trim = [s for s in non_selectionnes if s['freq']=='trimestriel']
    reste_mens = [s for s in non_selectionnes if s['freq']=='mensuel']
    deflate_branche = [s for s in nouvelles if s['branche']==prefix]

    ws = wb.create_sheet(nom_feuille[:31])
    ws.sheet_view.showGridLines = False
    col = 1
    col = ecrire_bloc(ws, col, "CIBLE", FILL_CIBLE, construire_dict([{'nom':f"VA {nom_feuille} (Mdh)",'serie':brut_cible}]))
    col = ecrire_bloc(ws, col, f"ANCRES - Trim. ({len(ancres_trim)})", FILL_ANCRE, construire_dict(ancres_trim))
    col = ecrire_bloc(ws, col, f"ANCRES - Mens. ({len(ancres_mens)})", FILL_ANCRE, construire_dict(ancres_mens))
    col = ecrire_bloc(ws, col, f"SÉLECTION N2 - Trim. ({len(sel_trim)})", FILL_SELECT, construire_dict(sel_trim))
    col = ecrire_bloc(ws, col, f"SÉLECTION N2 - Mens. ({len(sel_mens)})", FILL_SELECT, construire_dict(sel_mens))
    if deflate_branche:
        col = ecrire_bloc(ws, col, f"SÉLECTION APRÈS DÉFLATION ({len(deflate_branche)})", FILL_DEFLATE, construire_dict(deflate_branche))
    col = ecrire_bloc(ws, col, f"NON SÉLECTIONNÉ - Trim. ({len(reste_trim)})", FILL_T, construire_dict(reste_trim))
    col = ecrire_bloc(ws, col, f"NON SÉLECTIONNÉ - Mens. ({len(reste_mens)})", FILL_M, construire_dict(reste_mens))
    ws.freeze_panes = "A4"
    return dict(ancres=len(ancres_objs), select=len(selectionnes), deflate=len(deflate_branche), non_select=len(non_selectionnes))

for nom_feuille, prefix in BRANCHES_STANDARD:
    with open(f'{DOSSIER}/data/{prefix}_brut.pkl', 'rb') as f:
        brut = pickle.load(f)
    d = res[prefix]
    s = traiter_branche(nom_feuille, prefix, d['ancres_objs'], d['res_reste'], brut['cible_retropolee'])
    recap.append((nom_feuille, s['ancres'], s['select'], s['deflate'], s['non_select']))
    print(f"{nom_feuille:28s} : ancres={s['ancres']} select_N2={s['select']} deflate={s['deflate']} non_select={s['non_select']}")

with open(f'{DOSSIER}/data/immobilier_brut.pkl', 'rb') as f:
    immo_brut = pickle.load(f)
s = traiter_branche('Immobilier', 'immobilier', res_immo['ancres_objs'], res_immo['res_reste'], immo_brut['cible_retropolee'])
recap.append(('Immobilier', s['ancres'], s['select'], s['deflate'], s['non_select']))
print(f"{'Immobilier':28s} : ancres={s['ancres']} select_N2={s['select']} deflate={s['deflate']} non_select={s['non_select']}")

ws = wb.create_sheet('Sommaire', 0)
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:F1")
ws["A1"] = "GDPNow-Maroc — Classeur complet, avec sélection après déflation"
ws["A1"].font = Font(size=13, bold=True, color="FFFFFF"); ws["A1"].fill = FILL_DEFLATE
ws["A1"].alignment = Alignment(vertical="center", indent=1); ws.row_dimensions[1].height = 22

headers = ["Branche", "Ancres (N1)", "Sélection N2", "Sélection après déflation", "Non sélectionné", "Feuille"]
for j, h in enumerate(headers):
    c = ws.cell(row=3, column=j+1, value=h)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_DEFLATE; c.border = BORDER
for col_l, w in zip("ABCDEF", [26,12,14,20,16,22]):
    ws.column_dimensions[col_l].width = w

ta=ts=td=tn=0
for i, (nom, na, ns, nd, nn) in enumerate(recap):
    r=4+i
    ws.cell(row=r,column=1,value=nom).border=BORDER
    ws.cell(row=r,column=2,value=na).border=BORDER
    ws.cell(row=r,column=3,value=ns).border=BORDER
    ws.cell(row=r,column=4,value=nd).border=BORDER
    ws.cell(row=r,column=5,value=nn).border=BORDER
    ws.cell(row=r,column=6,value=nom[:31]).border=BORDER
    ta+=na; ts+=ns; td+=nd; tn+=nn
r=4+len(recap)
for c_idx,val in zip([1,2,3,4,5],["TOTAL",ta,ts,td,tn]):
    cc=ws.cell(row=r,column=c_idx,value=val); cc.font=Font(bold=True); cc.border=BORDER

r2=r+2
ws.cell(row=r2,column=1,value="Légende : vert=cible | or=ANCRES (jamais testées) | violet=SÉLECTION N2 | sarcelle=SÉLECTION APRÈS DÉFLATION | bleu/orange=non sélectionné")
ws.cell(row=r2+1,column=1,value=f"Verification : {ta} ancres + {ts} select N2 + {td} deflation + {tn} non-select = {ta+ts+td+tn} (dont 5 recuperees uniquement grace a la deflation)")

wb.save('classeur_complet_avec_deflation.xlsx')
print(f"\nTOTAL: {ta}+{ts}+{td}+{tn} = {ta+ts+td+tn}")
