# -*- coding: utf-8 -*-
"""
03_extraction_selection_peche.py -- Phase 2 : extraction et selection pour
Peche, avec lecture optimisee (bloc, pas cellule par cellule).
"""
from openpyxl import load_workbook
from datetime import date
import re, pickle
import numpy as np
import pandas as pd
from scipy import stats
from collections import Counter

FICHIER = '/mnt/user-data/uploads/Etude_sectorielle_Maroc_2_complete.xlsx'
BRANCHE = 'Pêche'
TODAY = date(2026, 8, 19)
DATE_DEBUT_ESTIMATION = date(2010, 1, 1)
DATE_FIN_TRAIN = date(2021, 1, 1)

def parse_date_cell(v):
    if v is None: return None
    if hasattr(v, 'year'): return date(v.year, v.month, 1)
    s = str(v).strip()
    m = re.match(r'T(\d)[\-\s](\d{4})', s)
    if m: t,y=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})T(\d)', s)
    if m: y,t=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})M(\d{1,2})', s)
    if m: y,mo=int(m.group(1)),int(m.group(2)); return date(y,mo,1)
    mois_fr={'janv':1,'févr':2,'fevr':2,'mars':3,'avr':4,'mai':5,'juin':6,'juil':7,'août':8,'aout':8,'sept':9,'oct':10,'nov':11,'déc':12,'dec':12}
    m = re.match(r'([a-zéû]+)\.?-(\d{2,4})', s.lower())
    if m:
        mo_str,y=m.group(1),int(m.group(2))
        if y<100: y+=2000 if y<50 else 1900
        if mo_str in mois_fr: return date(y,mois_fr[mo_str],1)
    m = re.match(r'^(\d{4})$', s)
    if m: return date(int(m.group(1)),1,1)
    return None

wb = load_workbook(FICHIER, read_only=True, data_only=True)
ws = wb[BRANCHE]
max_r, max_c = ws.max_row, ws.max_column
print(f"Lecture en bloc : {max_r} lignes x {max_c} colonnes...")

# Lecture complete en une seule passe (rapide en read_only via iter_rows)
grille = [[None]*(max_c+1) for _ in range(max_r+1)]
for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_r, values_only=True), start=1):
    for c_idx, val in enumerate(row, start=1):
        grille[r_idx][c_idx] = val
print("Lecture terminee.")

DATE_MARKERS = {'Mois','Trimestre','Année'}
col_dates = {}
for c in range(1, max_c+1):
    v3, v4 = grille[3][c], grille[4][c]
    if v4 in DATE_MARKERS: col_dates[c]=v4
    elif v3 in DATE_MARKERS: col_dates[c]=v3
date_cols_sorted = sorted(col_dates.keys())
blocs = []
for i,dc in enumerate(date_cols_sorted):
    fin = date_cols_sorted[i+1] if i+1<len(date_cols_sorted) else max_c+1
    blocs.append((dc, [c for c in range(dc,fin) if c!=dc]))

def extraire_serie(date_col, val_col, ligne_debut=5):
    serie = []
    vides = 0
    for r in range(ligne_debut, max_r+1):
        d = parse_date_cell(grille[r][date_col])
        v = grille[r][val_col]
        if d is None and v is None:
            vides += 1
            if vides >= 30: break
        else:
            vides = 0
            if d is not None and isinstance(v,(int,float)):
                serie.append((d,float(v)))
    return serie

cible_base = extraire_serie(1,2)
cible_retro = extraire_serie(12,13)
print(f"Cible base 2014: {len(cible_base)} obs | retropolee: {len(cible_retro)} obs")

COLS_CIBLE_VARIANTES = set(range(96, max_c+1))
indicateurs = []
for dc, val_cols in blocs:
    if dc in (1,12): continue
    for vc in val_cols:
        if vc in COLS_CIBLE_VARIANTES: continue
        nom = grille[3][vc]
        source = grille[4][vc]
        if not nom: continue
        serie = extraire_serie(dc,vc)
        if len(serie) < 4: continue
        indicateurs.append(dict(nom=str(nom), source=str(source) if source else '', serie=serie, colonne=vc))

print(f"{len(indicateurs)} indicateurs bruts extraits pour {BRANCHE}")
with open('/home/claude/phase2/data/peche_brut.pkl','wb') as f:
    pickle.dump(dict(cible_base=cible_base, cible_retropolee=cible_retro, indicateurs=indicateurs), f)

def categorie_economique(nom, source):
    src = source.lower()
    if 'débarquement' in src and 'espèce' in src: return 'especes'
    if 'débarquement' in src and 'port' in src: return 'ports'
    if 'comptes débiteurs' in nom.lower() or "crédits à l'équipement" in nom.lower(): return 'credit_combine'
    return None

candidats_eco = []
for ind in indicateurs:
    cat = categorie_economique(ind['nom'], ind['source'])
    if cat is None: continue
    dates_only = sorted(d for d,v in ind['serie'])
    ecarts = [(dates_only[i+1]-dates_only[i]).days for i in range(len(dates_only)-1)]
    med = np.median(ecarts) if ecarts else 999
    freq = "mensuel" if med<=40 else ("trimestriel" if med<=100 else "annuel")
    if freq == "annuel": continue
    candidats_eco.append(dict(**ind, freq=freq, categorie=cat))

print(f"\nCandidats apres filtre economique + frequence : {len(candidats_eco)}")
print("Par categorie:", Counter(c['categorie'] for c in candidats_eco))

SEUIL_LONGUEUR_MIN, SEUIL_FRAICHEUR_JOURS, SEUIL_DENSITE_MIN = 12, 450, 0.80
def expected_count(dmin,dmax,freq):
    months = (dmax.year-dmin.year)*12+(dmax.month-dmin.month)+1
    return months if freq=="mensuel" else months/3

candidats_qualite, rejets_qualite = [], []
for c in candidats_eco:
    dates_only = sorted(d for d,v in c['serie'])
    dmin,dmax = dates_only[0], dates_only[-1]
    if len(c['serie']) < SEUIL_LONGUEUR_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],'C2-longueur')); continue
    if (TODAY-dmax).days > SEUIL_FRAICHEUR_JOURS:
        rejets_qualite.append((c['nom'],c['colonne'],f'C3-fraicheur({dmax})')); continue
    exp = expected_count(dmin,dmax,c['freq'])
    densite = len(c['serie'])/exp if exp>0 else 0
    if densite < SEUIL_DENSITE_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],f'C4-densite({densite:.0%})')); continue
    candidats_qualite.append(dict(**c, densite=densite, debut=dmin, fin=dmax))

print(f"\nCandidats apres filtre de qualite : {len(candidats_qualite)}  (rejetes: {len(rejets_qualite)})")
for nom,col,code in rejets_qualite[:20]:
    print(f"  [rejete] col{col} {nom[:40]:40s} -- {code}")

def log_diff(serie):
    s = pd.Series({d:v for d,v in serie}).sort_index(); s=s[s>0]
    return np.log(s).diff().dropna()
def to_quarterly(serie):
    s = pd.Series({d:v for d,v in serie}).sort_index()
    s.index = pd.to_datetime(s.index)
    q = s.resample("QS").mean().dropna()
    return [(d.date(),v) for d,v in q.items()]

cible = sorted(cible_retro)
cible_ld_full = log_diff(cible)
cible_ld_train = cible_ld_full[(cible_ld_full.index>=DATE_DEBUT_ESTIMATION)&(cible_ld_full.index<=DATE_FIN_TRAIN)]

SEUIL_ALPHA, SEUIL_R_PLANCHER, NB_MIN_TRIM = 0.10, 0.15, 12
resultats_test = []
for c in candidats_qualite:
    serie_q = to_quarterly(c['serie']) if c['freq']=='mensuel' else c['serie']
    ind_ld = log_diff(serie_q)
    ind_ld_train = ind_ld[(ind_ld.index>=DATE_DEBUT_ESTIMATION)&(ind_ld.index<=DATE_FIN_TRAIN)]
    common = cible_ld_train.index.intersection(ind_ld_train.index)
    n = len(common)
    if n < NB_MIN_TRIM:
        resultats_test.append(dict(nom=c['nom'],colonne=c['colonne'],categorie=c['categorie'],r=None,p=None,n=n,retenu=False,motif=f"n={n}<{NB_MIN_TRIM}"))
        continue
    r = np.corrcoef(cible_ld_train.loc[common], ind_ld_train.loc[common])[0,1]
    if np.isnan(r):
        resultats_test.append(dict(nom=c['nom'],colonne=c['colonne'],categorie=c['categorie'],r=None,p=None,n=n,retenu=False,motif="r indefini")); continue
    t = r*np.sqrt((n-2)/(1-r**2)) if abs(r)<1 else np.inf
    p = 2*(1-stats.t.cdf(abs(t), df=n-2))
    retenu = (p<SEUIL_ALPHA) and (abs(r)>=SEUIL_R_PLANCHER)
    resultats_test.append(dict(nom=c['nom'],colonne=c['colonne'],categorie=c['categorie'],r=round(r,3),p=round(p,3),n=n,retenu=retenu,
                                 motif="RETENU" if retenu else ("sous plancher" if abs(r)<SEUIL_R_PLANCHER else "non signif.")))

retenus = [r for r in resultats_test if r['retenu']]
print(f"\n=== RESULTAT DU TEST (train {DATE_DEBUT_ESTIMATION} a {DATE_FIN_TRAIN}) ===")
print(f"TOTAL RETENU POUR PECHE : {len(retenus)} / {len(resultats_test)} testes\n")
for r in sorted(retenus, key=lambda x:-abs(x['r'])):
    print(f"  col{r['colonne']:3d} [{r['categorie']:14s}] {r['nom'][:35]:35s} r={r['r']:+.3f} p={r['p']:.3f} n={r['n']}")

with open('/home/claude/phase2/resultats/peche_selection.pkl','wb') as f:
    pickle.dump(dict(candidats_qualite=candidats_qualite, resultats_test=resultats_test, rejets_qualite=rejets_qualite), f)
