# -*- coding: utf-8 -*-
"""
04_extraction_filtre_industrie_transfo.py -- Phase 2 : extraction, filtre
economique, filtre de qualite pour Industrie de transformation.
ARRET AU FILTRE DE QUALITE (pas de test statistique a ce stade).
"""
from openpyxl import load_workbook
from datetime import date
import re, pickle
import numpy as np
from collections import Counter

FICHIER = '/mnt/user-data/uploads/Etude_sectorielle_Maroc_2_complete.xlsx'
BRANCHE = 'Industrie de transformation'
TODAY = date(2026, 8, 19)

def parse_date_cell(v):
    if v is None: return None
    if hasattr(v, 'year'): return date(v.year, v.month, 1)
    s = str(v).strip()
    m = re.match(r'T(\d)[\-\s](\d{4})', s)
    if m: t,y=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})T(\d)', s)
    if m: y,t=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})M(\d{1,2})', s)
    if m: y,mo=int(m.group(1)),int(m.group(2)); return date(y,mo,1)
    mois_fr={'janv':1,'févr':2,'fevr':2,'mars':3,'avr':4,'mai':5,'juin':6,'juil':7,'août':8,'aout':8,'sept':9,'oct':10,'nov':11,'déc':12,'dec':12}
    m = re.match(r'([a-zéû]+)\.?-(\d{2,4})', s.lower())
    if m:
        mo_str,y=m.group(1),int(m.group(2))
        if y<100: y+=2000 if y<50 else 1900
        if mo_str in mois_fr: return date(y,mois_fr[mo_str],1)
    m = re.match(r'^(\d{4})$', s)
    if m: return date(int(m.group(1)),1,1)
    return None

wb = load_workbook(FICHIER, read_only=True, data_only=True)
ws = wb[BRANCHE]
max_r, max_c = ws.max_row, ws.max_column
print(f"Lecture en bloc : {max_r} lignes x {max_c} colonnes...")
grille = [[None]*(max_c+1) for _ in range(max_r+1)]
for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_r, values_only=True), start=1):
    for c_idx, val in enumerate(row, start=1):
        grille[r_idx][c_idx] = val
print("Lecture terminee.")

DATE_MARKERS = {'Mois','Trimestre','Année'}
col_dates = {}
for c in range(1, max_c+1):
    v3, v4 = grille[3][c], grille[4][c]
    if v4 in DATE_MARKERS: col_dates[c]=v4
    elif v3 in DATE_MARKERS: col_dates[c]=v3
date_cols_sorted = sorted(col_dates.keys())
blocs = []
for i,dc in enumerate(date_cols_sorted):
    fin = date_cols_sorted[i+1] if i+1<len(date_cols_sorted) else max_c+1
    blocs.append((dc, [c for c in range(dc,fin) if c!=dc]))

def extraire_serie(date_col, val_col, ligne_debut=5):
    serie = []; vides = 0
    for r in range(ligne_debut, max_r+1):
        d = parse_date_cell(grille[r][date_col])
        v = grille[r][val_col]
        if d is None and v is None:
            vides += 1
            if vides >= 30: break
        else:
            vides = 0
            if d is not None and isinstance(v,(int,float)): serie.append((d,float(v)))
    return serie

cible_retro = extraire_serie(48, 49)
print(f"Cible retropolee : {len(cible_retro)} obs")

# Colonnes de variantes de la cible (109+) et colonnes confirmees doublons
COLS_CIBLE_VARIANTES = set(range(109, max_c+1))
COLS_DOUBLONS_CONFIRMES = {102, 104}  # = col43, col44 (verifie chiffre pour chiffre)

# Sous-secteurs inferes pour les 5 paires Comptes debiteurs / Credits equipement
# (col37-46), par la position dans la table source officielle Bank Al-Maghrib
# (meme ordre que les 5 colonnes de credit bancaire, cols 8-12) :
SOUS_SECTEURS_INFERES = {
    37: "alimentaire et tabac", 38: "alimentaire et tabac",
    39: "textile-habillement-cuir", 40: "textile-habillement-cuir",
    41: "chimique-parachimique", 42: "chimique-parachimique",
    43: "métallurgique-mécanique-électrique", 44: "métallurgique-mécanique-électrique",
    45: "manufacturières diverses", 46: "manufacturières diverses",
}

indicateurs = []
for dc, val_cols in blocs:
    if dc in (1, 48): continue
    for vc in val_cols:
        if vc in COLS_CIBLE_VARIANTES or vc in COLS_DOUBLONS_CONFIRMES: continue
        nom = grille[3][vc]
        source = grille[4][vc]
        if not nom: continue
        serie = extraire_serie(dc,vc)
        if len(serie) < 4: continue
        indicateurs.append(dict(nom=str(nom), source=str(source) if source else '', serie=serie, colonne=vc))

print(f"{len(indicateurs)} indicateurs bruts extraits (apres retrait des 2 doublons confirmes)")

with open('/home/claude/phase2/data/industrie_transfo_brut.pkl','wb') as f:
    pickle.dump(dict(cible_retropolee=cible_retro, indicateurs=indicateurs), f)

# ============================================================================
# FILTRE ECONOMIQUE A PRIORI -- par categorie
# ============================================================================
def categorie_economique(nom, source, colonne):
    n, s = nom.lower(), source.lower()
    if 'taux d' in n and 'utilisation' in n and 'capacit' in n: return 'TUC'
    if n.startswith('ipi'): return 'IPI'
    if 'crédit bancaire' in n: return 'credit_bancaire_sectoriel'
    if colonne in SOUS_SECTEURS_INFERES: return 'credit_ventilation_inferee'
    if 'enquête' in s or 'conjoncture' in s: return 'enquete_conjoncture'
    if colonne in (99,101,103): return 'credit_agregat_secteur'
    if colonne == 100: return 'credit_a_terme_metallurgique'
    if 'produits finis' in n or 'demi produits' in n or 'demi-produits' in n: return 'commerce_exterieur_intrants'
    if 'utilisation des capacités' in n: return 'TUC'
    return None

candidats_eco = []
for ind in indicateurs:
    cat = categorie_economique(ind['nom'], ind['source'], ind['colonne'])
    if cat is None: continue
    dates_only = sorted(d for d,v in ind['serie'])
    ecarts = [(dates_only[i+1]-dates_only[i]).days for i in range(len(dates_only)-1)]
    med = np.median(ecarts) if ecarts else 999
    freq = "mensuel" if med<=40 else ("trimestriel" if med<=100 else "annuel")
    if freq == "annuel": continue
    candidats_eco.append(dict(**ind, freq=freq, categorie=cat))

print(f"\nCandidats apres filtre economique + frequence : {len(candidats_eco)}")
print("Par categorie:", Counter(c['categorie'] for c in candidats_eco))

# ============================================================================
# FILTRE DE QUALITE
# ============================================================================
SEUIL_LONGUEUR_MIN, SEUIL_FRAICHEUR_JOURS, SEUIL_DENSITE_MIN = 12, 450, 0.80
def expected_count(dmin,dmax,freq):
    months = (dmax.year-dmin.year)*12+(dmax.month-dmin.month)+1
    return months if freq=="mensuel" else months/3

candidats_qualite, rejets_qualite = [], []
for c in candidats_eco:
    dates_only = sorted(d for d,v in c['serie'])
    dmin,dmax = dates_only[0], dates_only[-1]
    if len(c['serie']) < SEUIL_LONGUEUR_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],c['categorie'],'C2-longueur')); continue
    if (TODAY-dmax).days > SEUIL_FRAICHEUR_JOURS:
        rejets_qualite.append((c['nom'],c['colonne'],c['categorie'],f'C3-fraicheur({dmax})')); continue
    exp = expected_count(dmin,dmax,c['freq'])
    densite = len(c['serie'])/exp if exp>0 else 0
    if densite < SEUIL_DENSITE_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],c['categorie'],f'C4-densite({densite:.0%})')); continue
    candidats_qualite.append(dict(**c, densite=densite, debut=dmin, fin=dmax))

print(f"\nCandidats apres filtre de qualite : {len(candidats_qualite)}  (rejetes: {len(rejets_qualite)})")
print("Par categorie (retenus):", Counter(c['categorie'] for c in candidats_qualite))
print("\nRejets:")
for nom,col,cat,code in rejets_qualite:
    print(f"  col{col:3d} [{cat:30s}] {nom[:45]:45s} -- {code}")

with open('/home/claude/phase2/resultats/industrie_transfo_selection.pkl','wb') as f:
    pickle.dump(dict(candidats_qualite=candidats_qualite, rejets_qualite=rejets_qualite), f)

import pandas as pd
rows = [dict(colonne=c['colonne'],nom=c['nom'],categorie=c['categorie'],freq=c['freq'],source=c['source'],
             nobs=len(c['serie']),debut=c['debut'],fin=c['fin'],densite=round(c['densite'],3)) for c in candidats_qualite]
pd.DataFrame(rows).sort_values(['categorie','nom']).to_csv(
    '/home/claude/phase2/resultats/industrie_transfo_apres_filtre_qualite.csv', index=False, encoding='utf-8-sig')
print(f"\nCSV exporte : {len(rows)} lignes")
