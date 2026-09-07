# -*- coding: utf-8 -*-
"""09_extraction_filtre_construction.py -- Phase 2, Construction, sans fraicheur."""
from openpyxl import load_workbook
from datetime import date
import re, pickle
import numpy as np
from collections import Counter

FICHIER = '/mnt/user-data/uploads/Etude_sectorielle_Maroc_2_complete.xlsx'
BRANCHE = 'Construction'

def parse_date_cell(v):
    if v is None: return None
    if hasattr(v, 'year'): return date(v.year, v.month, 1)
    s = str(v).strip()
    m = re.match(r'T(\d)[\-\s](\d{4})', s)
    if m: t,y=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})T(\d)', s)
    if m: y,t=int(m.group(1)),int(m.group(2)); return date(y,(t-1)*3+1,1)
    m = re.match(r'(\d{4})M(\d{1,2})', s)
    if m: y,mo=int(m.group(1)),int(m.group(2)); return date(y,mo,1)
    mois_fr={'janv':1,'févr':2,'fevr':2,'mars':3,'avr':4,'mai':5,'juin':6,'juil':7,'août':8,'aout':8,'sept':9,'oct':10,'nov':11,'déc':12,'dec':12}
    m = re.match(r'([a-zéû]+)\.?-(\d{2,4})', s.lower())
    if m:
        mo_str,y=m.group(1),int(m.group(2))
        if y<100: y+=2000 if y<50 else 1900
        if mo_str in mois_fr: return date(y,mois_fr[mo_str],1)
    m = re.match(r'^(\d{4})$', s)
    if m: return date(int(m.group(1)),1,1)
    return None

wb = load_workbook(FICHIER, read_only=True, data_only=True)
ws = wb[BRANCHE]
max_r, max_c = ws.max_row, ws.max_column
grille = [[None]*(max_c+1) for _ in range(max_r+1)]
for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_r, values_only=True), start=1):
    for c_idx, val in enumerate(row, start=1):
        grille[r_idx][c_idx] = val

DATE_MARKERS = {'Mois','Trimestre','Année'}
col_dates = {}
for c in range(1, max_c+1):
    v3, v4 = grille[3][c], grille[4][c]
    if v4 in DATE_MARKERS: col_dates[c]=v4
    elif v3 in DATE_MARKERS: col_dates[c]=v3
date_cols_sorted = sorted(col_dates.keys())
blocs = []
for i,dc in enumerate(date_cols_sorted):
    fin = date_cols_sorted[i+1] if i+1<len(date_cols_sorted) else max_c+1
    blocs.append((dc, [c for c in range(dc,fin) if c!=dc]))

def extraire_serie(date_col, val_col, ligne_debut=5):
    serie = []; vides = 0
    for r in range(ligne_debut, max_r+1):
        d = parse_date_cell(grille[r][date_col])
        v = grille[r][val_col]
        if d is None and v is None:
            vides += 1
            if vides >= 30: break
        else:
            vides = 0
            if d is not None and isinstance(v,(int,float)): serie.append((d,float(v)))
    return serie

cible_retro = extraire_serie(12, 13)
print(f"Cible retropolee : {len(cible_retro)} obs")

COLS_CIBLE_VARIANTES = set(range(50, 60))
COLS_CUMULE_EXCLUES = {17}
COLS_VIDES = {34, 35}  # confirme vides (None) a l'inspection

indicateurs = []
for dc, val_cols in blocs:
    if dc in (1,12): continue
    for vc in val_cols:
        if vc in COLS_CIBLE_VARIANTES or vc in COLS_CUMULE_EXCLUES or vc in COLS_VIDES:
            continue
        nom = grille[3][vc]; source = grille[4][vc]
        if not nom: continue
        serie = extraire_serie(dc,vc)
        if len(serie) < 4: continue
        indicateurs.append(dict(nom=str(nom), source=str(source) if source else '', serie=serie, colonne=vc))

print(f"{len(indicateurs)} indicateurs bruts extraits")
with open('/home/claude/phase2/data/construction_brut.pkl','wb') as f:
    pickle.dump(dict(cible_retropolee=cible_retro, indicateurs=indicateurs), f)

def categorie_economique(nom, source, colonne):
    n, s = nom.lower(), source.lower()
    if 'crédit bancaire' in n or 'comptes débiteurs' in n or "crédits à l'équipement" in n:
        return 'credit_sectoriel'
    if colonne == 5: return 'ciment_national'
    if colonne in range(18,34) or colonne == 36 or colonne in range(37,49):
        return 'ciment_regional'
    if 'pierres pour la construction' in n: return 'IPM_materiaux'
    return None

candidats_eco = []
for ind in indicateurs:
    cat = categorie_economique(ind['nom'], ind['source'], ind['colonne'])
    if cat is None: continue
    dates_only = sorted(d for d,v in ind['serie'])
    ecarts = [(dates_only[i+1]-dates_only[i]).days for i in range(len(dates_only)-1)]
    med = np.median(ecarts) if ecarts else 999
    freq = "mensuel" if med<=40 else ("trimestriel" if med<=100 else "annuel")
    if freq == "annuel": continue
    candidats_eco.append(dict(**ind, freq=freq, categorie=cat))

print(f"\nCandidats apres filtre economique + frequence : {len(candidats_eco)}")
print("Par categorie:", Counter(c['categorie'] for c in candidats_eco))

SEUIL_LONGUEUR_MIN, SEUIL_DENSITE_MIN = 12, 0.80
def expected_count(dmin,dmax,freq):
    months = (dmax.year-dmin.year)*12+(dmax.month-dmin.month)+1
    return months if freq=="mensuel" else months/3

candidats_qualite, rejets_qualite = [], []
for c in candidats_eco:
    dates_only = sorted(d for d,v in c['serie'])
    dmin,dmax = dates_only[0], dates_only[-1]
    if len(c['serie']) < SEUIL_LONGUEUR_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],c['categorie'],'C2-longueur')); continue
    exp = expected_count(dmin,dmax,c['freq'])
    densite = len(c['serie'])/exp if exp>0 else 0
    if densite < SEUIL_DENSITE_MIN:
        rejets_qualite.append((c['nom'],c['colonne'],c['categorie'],f'C4-densite({densite:.0%})')); continue
    candidats_qualite.append(dict(**c, densite=densite, debut=dmin, fin=dmax))

print(f"\nCandidats apres filtre de qualite (sans fraicheur) : {len(candidats_qualite)}  (rejetes: {len(rejets_qualite)})")
print("Par categorie (retenus):", Counter(c['categorie'] for c in candidats_qualite))
print("\nRejets:")
for nom,col,cat,code in rejets_qualite:
    print(f"  col{col:3d} [{cat:18s}] {nom[:40]:40s} -- {code}")
print("\nRetenus:")
for c in candidats_qualite:
    print(f"  col{c['colonne']:3d} [{c['categorie']:18s}] {c['nom'][:45]:45s} {c['freq']:10s} {c['debut']} a {c['fin']}")

with open('/home/claude/phase2/resultats/construction_sans_fraicheur.pkl','wb') as f:
    pickle.dump(candidats_qualite, f)
