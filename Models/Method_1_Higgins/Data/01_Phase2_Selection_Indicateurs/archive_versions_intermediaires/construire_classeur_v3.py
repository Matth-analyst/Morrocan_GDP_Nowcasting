# -*- coding: utf-8 -*-
"""
construire_classeur_v3.py -- Classeur final, structure fidele au classeur
maitre : blocs cote a cote (cible | trimestriel | mensuel), chaque bloc
avec SA PROPRE colonne de dates juste avant lui.
"""
import pickle
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

DOSSIER = '/home/claude/phase2'
FILL_CIBLE = PatternFill("solid", fgColor="548235")   # cible : vert
FILL_T = PatternFill("solid", fgColor="1F4E78")        # trimestriel : bleu fonce
FILL_M = PatternFill("solid", fgColor="C55A11")        # mensuel : orange
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SEPARATEUR_LARGEUR = 2   # colonnes vides entre chaque bloc

wb = Workbook()
wb.remove(wb.active)

def ecrire_bloc(ws, col_depart, titre, fill, series_dict):
    """Ecrit un bloc [Date | col1 | col2 | ...] a partir de col_depart.
    series_dict : {nom_colonne: {date: valeur}}. Renvoie la colonne
    suivante disponible (apres le separateur)."""
    if not series_dict:
        return col_depart
    df = pd.DataFrame(series_dict).sort_index()
    n_col = len(df.columns)

    # Titre du bloc (ligne 1)
    c1 = get_column_letter(col_depart)
    c2 = get_column_letter(col_depart + n_col)
    ws.merge_cells(f"{c1}1:{c2}1")
    cell = ws.cell(row=1, column=col_depart, value=titre)
    cell.font = Font(size=11, bold=True, color="FFFFFF")
    cell.fill = fill
    cell.alignment = Alignment(vertical="center", horizontal="center")

    # En-tetes (ligne 3) : Date, puis chaque serie
    cell = ws.cell(row=3, column=col_depart, value="Date")
    cell.font = Font(bold=True, color="FFFFFF"); cell.fill = fill; cell.border = BORDER
    ws.column_dimensions[get_column_letter(col_depart)].width = 12

    for j, col in enumerate(df.columns):
        cc = col_depart + 1 + j
        cell = ws.cell(row=3, column=cc, value=col)
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(cc)].width = 15
    ws.row_dimensions[3].height = 60

    # Donnees a partir de la ligne 4
    for i, (d, row) in enumerate(df.iterrows()):
        r = 4 + i
        c0 = ws.cell(row=r, column=col_depart, value=d)
        c0.number_format = "YYYY-MM-DD"; c0.border = BORDER; c0.font = Font(size=9)
        for j, col in enumerate(df.columns):
            cc = col_depart + 1 + j
            val = row[col]
            cell = ws.cell(row=r, column=cc,
                            value=round(float(val), 4) if pd.notna(val) else None)
            if pd.notna(val): cell.number_format = "#,##0.0000"
            cell.border = BORDER; cell.font = Font(size=9)

    return col_depart + 1 + n_col + SEPARATEUR_LARGEUR

def construire_dict_series(cible_nom, cible_serie, indicateurs):
    d = {}
    if cible_nom is not None:
        d[cible_nom] = dict(cible_serie)
    for ind in indicateurs:
        nom = ind['nom']; base = nom; i = 2
        while nom in d:
            nom = f"{base} ({i})"; i += 1
        d[nom] = dict(ind['serie'])
    return d

# ============================================================================
BRANCHES_STANDARD = [
    ('Agriculture', 'agriculture'), ('Peche', 'peche'),
    ('Industrie transformation', 'industrie_transfo'), ('Industrie extraction', 'ind_extraction'),
    ('Finances assurances', 'finances'), ('Hebergement restauration', 'hebergement'),
    ('Construction', 'construction'), ('Commerce', 'commerce'),
    ('Transports', 'transports'), ('Electricite gaz eau', 'electricite'),
    ('Information communication', 'infocom'),
]
recap = []

for nom_feuille, prefix in BRANCHES_STANDARD:
    with open(f'{DOSSIER}/data/{prefix}_brut.pkl', 'rb') as f:
        brut = pickle.load(f)
    with open(f'{DOSSIER}/resultats/{prefix}_sans_fraicheur.pkl', 'rb') as f:
        indicateurs = pickle.load(f)
    ind_trim = [i for i in indicateurs if i['freq']=='trimestriel']
    ind_mens = [i for i in indicateurs if i['freq']=='mensuel']

    ws = wb.create_sheet(nom_feuille[:31])
    ws.sheet_view.showGridLines = False

    col = 1
    col = ecrire_bloc(ws, col, "CIBLE", FILL_CIBLE,
                        construire_dict_series(f"VA {nom_feuille} (Mdh)", brut['cible_retropolee'], []))
    col = ecrire_bloc(ws, col, f"TRIMESTRIEL ({len(ind_trim)})", FILL_T,
                        construire_dict_series(None, None, ind_trim))
    col = ecrire_bloc(ws, col, f"MENSUEL ({len(ind_mens)})", FILL_M,
                        construire_dict_series(None, None, ind_mens))

    ws.freeze_panes = "A4"
    recap.append((nom_feuille, len(ind_trim), len(ind_mens)))
    print(f"{nom_feuille:30s} : {len(ind_trim):3d} trim. + {len(ind_mens):3d} mens.")

# ============================================================================
# Immobilier
# ============================================================================
with open(f'{DOSSIER}/data/immobilier_brut.pkl', 'rb') as f:
    immo_brut = pickle.load(f)
COLS_CREDIT_IMMO = {5, 6, 71, 72}
credit_series = [ind for ind in immo_brut['indicateurs'] if ind['colonne'] in COLS_CREDIT_IMMO]

ipai_etendu = pd.read_csv(f'{DOSSIER.replace("/phase2","")}/ipai_nouveaux/ipai_variations_propre_etendue.csv',
                            sep=';', encoding='utf-8-sig')
def trim_to_date(t):
    trim, an = t.split('-')
    return date(int(an), (int(trim[1])-1)*3+1, 1)
IPAI_RETENUES = [
    ("Professionnel","prix"), ("Foncier","prix"), ("Bureau","prix"),
    ("Résidentiel","transactions"), ("Local commercial","transactions"),
    ("Local commercial","prix"), ("Professionnel","transactions"),
    ("Foncier","transactions"), ("Bureau","transactions"),
]
ipai_indicateurs = []
for cat, ind in IPAI_RETENUES:
    sub = ipai_etendu[(ipai_etendu['categorie']==cat) & (ipai_etendu['indicateur']==ind)]
    serie = [(trim_to_date(row['trimestre']), row['variation_trimestrielle']) for _, row in sub.iterrows()]
    ipai_indicateurs.append(dict(nom=f"IPAI {cat} — {ind} (var. trim. %)", serie=serie))
for ind in immo_brut['indicateurs']:
    if 'prose' in ind['nom'].lower():
        ipai_indicateurs.append(dict(nom="IPAI Global — prix (série 1998–2017, var. trim. %)", serie=ind['serie']))

ws = wb.create_sheet('Immobilier')
ws.sheet_view.showGridLines = False
col = 1
col = ecrire_bloc(ws, col, "CIBLE", FILL_CIBLE,
                    construire_dict_series('VA Immobilier (Mdh)', immo_brut['cible_retropolee'], []))
col = ecrire_bloc(ws, col, f"TRIMESTRIEL - IPAI ({len(ipai_indicateurs)})", FILL_T,
                    construire_dict_series(None, None, ipai_indicateurs))
col = ecrire_bloc(ws, col, f"MENSUEL - CREDIT ({len(credit_series)})", FILL_M,
                    construire_dict_series(None, None, credit_series))
ws.freeze_panes = "A4"
recap.append(('Immobilier', len(ipai_indicateurs), len(credit_series)))
print(f"{'Immobilier':30s} : {len(ipai_indicateurs):3d} trim. + {len(credit_series):3d} mens.")

# ============================================================================
# Sommaire
# ============================================================================
ws = wb.create_sheet('Sommaire', 0)
ws.sheet_view.showGridLines = False
ws.merge_cells("A1:E1")
ws["A1"] = "GDPNow-Maroc — Classeur des 434 séries retenues (Phase 2)"
ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
ws["A1"].fill = FILL_T
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 22

headers = ["Branche", "Trimestriel", "Mensuel", "Total", "Feuille"]
for j, h in enumerate(headers):
    c = ws.cell(row=3, column=j+1, value=h)
    c.font = Font(bold=True, color="FFFFFF"); c.fill = FILL_T; c.border = BORDER
for col_l, w in zip("ABCDE", [30,14,12,10,25]):
    ws.column_dimensions[col_l].width = w

total_t = total_m = 0
for i, (nom, nt, nm) in enumerate(recap):
    r = 4+i
    ws.cell(row=r, column=1, value=nom).border = BORDER
    ws.cell(row=r, column=2, value=nt).border = BORDER
    ws.cell(row=r, column=3, value=nm).border = BORDER
    ws.cell(row=r, column=4, value=nt+nm).border = BORDER
    ws.cell(row=r, column=5, value=nom[:31]).border = BORDER
    total_t += nt; total_m += nm

r = 4 + len(recap)
for c_idx, val in zip([1,2,3,4], ["TOTAL", total_t, total_m, total_t+total_m]):
    cc = ws.cell(row=r, column=c_idx, value=val)
    cc.font = Font(bold=True); cc.border = BORDER

r2 = r + 2
ws.cell(row=r2, column=1, value="Légende : vert = cible (VA) | bleu = bloc trimestriel | orange = bloc mensuel — chaque bloc a sa propre colonne de dates")
ws.cell(row=r2+1, column=1, value="Branches sans indicateur disponible (AR(4) pur) : Services aux entreprises, Administration publique, Éducation-santé, Autres services")

wb.save('/mnt/user-data/outputs/GDPNow_Maroc_series_retenues_Phase2.xlsx')
print(f"\nTotal : {total_t} trimestriels + {total_m} mensuels = {total_t+total_m}")
